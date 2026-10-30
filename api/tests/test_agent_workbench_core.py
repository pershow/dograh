import zipfile

import pytest

from api.schemas.agent_workbench import SimulationCase
from api.services.agent_workbench.errors import WorkbenchError
from api.services.agent_workbench.evaluation import (
    acceptable_improvement,
    check_assertions,
    resolve_node_references,
    tuning_evidence,
)
from api.services.agent_workbench.evaluation_policy import EvaluationPolicy
from api.services.agent_workbench.file_parser import parse_file


def test_node_assertions_resolve_ids_and_unique_names_without_false_passes():
    graph = {
        "nodes": [
            {"id": "10", "data": {"name": "成功结束"}},
            {"id": "11", "data": {"name": "未经过"}},
        ]
    }
    case = SimulationCase(
        name="routing",
        user_messages=["bye"],
        assertions=[
            {"kind": "node_visited", "value": "10"},
            {"kind": "node_visited", "value": "成功结束"},
            {"kind": "node_visited", "value": "未经过"},
        ],
    )
    references = resolve_node_references(case, graph)
    results = check_assertions(
        case,
        [],
        [
            {
                "type": "node_transition",
                "payload": {"node_id": "10", "node_name": "成功结束"},
            }
        ],
        {},
        True,
        [],
        references,
    )
    assert [r["passed"] for r in results] == [True, True, False]
    graph["nodes"][1]["data"]["name"] = "成功结束"
    with pytest.raises(WorkbenchError, match="unique node name"):
        resolve_node_references(case, graph)
    with pytest.raises(WorkbenchError, match="existing node ID"):
        resolve_node_references(case, {"nodes": []})


def test_tuning_evidence_preserves_failures_without_duplicate_event_transcripts():
    case = {
        "name": "routing",
        "status": "failed",
        "assertions": [{"kind": "node_visited", "value": "结束", "passed": False}],
        "transcript": [{"role": "assistant", "text": "A" * 30000}],
        "events": [
            {"type": "bot_speech", "payload": {"text": "DUPLICATED_SPEECH"}},
            {
                "type": "node_transition",
                "payload": {"node_id": "2", "node_name": "核实"},
            },
        ],
        "usage": {"prompt_tokens": 1000},
    }
    compact = tuning_evidence({"results": [case]})[0]
    assert compact["assertions"] == case["assertions"]
    assert compact["visited_nodes"][0]["node_id"] == "2"
    assert "events" not in compact and "usage" not in compact
    assert compact["transcript"][0]["text"].endswith("[truncated]")
    assert len(case["transcript"][0]["text"]) == 30000


def test_parser_documents_and_malformed_files(tmp_path):
    from docx import Document
    from openpyxl import Workbook

    doc = Document()
    doc.add_paragraph("Appointment rules")
    path = tmp_path / "rules.docx"
    doc.save(path)
    assert "Appointment rules" in parse_file(path)["text"]
    book = Workbook()
    book.active.append(["Price", "Free"])
    path = tmp_path / "prices.xlsx"
    book.save(path)
    assert "Free" in parse_file(path)["text"]
    path = tmp_path / "text.md"
    path.write_text("x" * 200001)
    assert parse_file(path)["truncated"]
    path = tmp_path / "bad.docx"
    with zipfile.ZipFile(path, "w") as z:
        z.writestr("word/vbaProject.bin", "malicious")
    with pytest.raises(ValueError, match="Macros"):
        parse_file(path)
    path = tmp_path / "binary.txt"
    path.write_bytes(b"\x00fake")
    with pytest.raises(ValueError, match="Binary"):
        parse_file(path)


def test_policy_requires_fixtures_before_any_live_effects():
    case = SimulationCase(
        name="call",
        user_messages=["hello"],
        assertions=[{"kind": "contains", "value": "free"}],
    )
    policy = EvaluationPolicy(case, None)
    for field, value in [
        ("tool_uuids", ["http-or-mcp"]),
        ("pre_call_fetch_url", "https://example.com"),
        ("document_uuids", ["doc"]),
        ("greeting_recording_id", 1),
    ]:
        with pytest.raises(WorkbenchError):
            policy.validate({"nodes": [{"type": "agentNode", "data": {field: value}}]})


async def test_policy_invokes_fixture_and_checks_arguments():
    case = SimulationCase(
        name="call",
        user_messages=["hello"],
        assertions=[{"kind": "tool_called", "value": "book"}],
        fixtures=[
            {
                "tool_uuid": "uuid",
                "name": "book",
                "parameters": {
                    "type": "object",
                    "properties": {"name": {"type": "string"}},
                    "required": ["name"],
                },
                "arguments": {"name": "Alice"},
                "result": {"booked": True},
            }
        ],
    )
    policy = EvaluationPolicy(case, None)
    policy.validate(
        {"nodes": [{"type": "agentNode", "data": {"tool_uuids": ["uuid"]}}]}
    )
    assert policy.invoke("book", {"name": "Alice"}) == {"booked": True}
    assert policy.invoke("book", {"name": "Bob"}) == {
        "error": "fixture_arguments_mismatch"
    }
    assert len(policy.calls) == 2 and policy.errors
    assert policy.calls[-1]["result"]["error"] == "fixture_arguments_mismatch"


def test_evaluation_cannot_pass_on_user_text_or_accept_regressions():
    case = SimulationCase(
        name="price",
        user_messages=["free"],
        assertions=[{"kind": "contains", "value": "free"}],
    )
    result = check_assertions(
        case,
        [{"role": "user", "text": "free"}, {"role": "assistant", "text": "paid"}],
        [],
        {},
        False,
        [],
    )
    assert not result[0]["passed"]
    old = {"results": [{"status": "passed"}, {"status": "failed"}]}
    assert acceptable_improvement(
        old, {"results": [{"status": "passed"}, {"status": "passed"}]}
    )
    for statuses in [
        ("failed", "passed"),
        ("passed", "inconclusive"),
        ("passed", "error"),
    ]:
        assert not acceptable_improvement(
            old, {"results": [{"status": s} for s in statuses]}
        )


async def test_runtime_policy_short_circuits_live_mcp_and_tool_resolution():
    from types import SimpleNamespace
    from unittest.mock import AsyncMock, Mock

    from api.services.workflow.pipecat_engine import PipecatEngine
    from api.services.workflow.pipecat_engine_custom_tools import CustomToolManager

    policy = Mock()
    policy.schemas.return_value = []
    engine = SimpleNamespace(
        evaluation_policy=policy,
        active_agent=Mock(),
        _get_organization_id=AsyncMock(
            side_effect=AssertionError("must not resolve live tools")
        ),
    )
    await PipecatEngine._open_mcp_sessions(engine)
    manager = CustomToolManager(engine)
    assert await manager.get_tool_schemas(["external"]) == []
    await manager.register_handlers(["external"])
    policy.register.assert_called_once()
    engine._get_organization_id.assert_not_called()


def test_node_assertion_reads_actual_runtime_event_payload():
    case = SimulationCase(
        name="visit",
        user_messages=["go"],
        assertions=[{"kind": "node_visited", "value": "node2"}],
    )
    assert check_assertions(
        case,
        [],
        [{"type": "node_transition", "payload": {"node_id": "node2"}}],
        {},
        False,
        [],
    )[0]["passed"]


def test_fixture_remote_schema_rejected():
    from pydantic import ValidationError

    from api.schemas.agent_workbench import ToolFixture

    with pytest.raises(ValidationError, match="remote resources"):
        ToolFixture(
            tool_uuid="tool",
            name="book",
            parameters={"$ref": "https://example.com/schema"},
        )


def test_qwen_compatibility_is_scoped_to_dashscope():
    from unittest.mock import patch

    from api.services.pipecat.service_factory import create_llm_service_from_provider

    with patch("api.services.pipecat.service_factory.OpenAILLMService") as service:
        create_llm_service_from_provider(
            "openai",
            "qwen3.6-27b",
            "test-key",
            base_url="https://dashscope.aliyuncs.com/compatible-mode/v1",
        )
        assert service.call_args.kwargs["settings"].extra == {
            "extra_body": {"enable_thinking": False}
        }
        assert service.return_value.supports_developer_role is False
    with patch("api.services.pipecat.service_factory.OpenAILLMService") as service:
        create_llm_service_from_provider(
            "openai", "gpt-4.1", "test-key", base_url="https://api.openai.com/v1"
        )
        assert service.call_args.kwargs["settings"].extra == {}


def test_image_parser_returns_bounded_model_input_without_source_metadata(tmp_path):
    from PIL import Image

    path = tmp_path / "sample.png"
    Image.new("RGB", (1600, 900), "white").save(path)
    parsed = parse_file(path)
    assert parsed["width"] == 1600 and parsed["model_image"].startswith(
        "data:image/jpeg;base64,"
    )
    assert len(parsed["model_image"]) < 50000


async def test_empty_turn_does_not_pass_using_the_fixed_greeting(monkeypatch):
    from types import SimpleNamespace
    from unittest.mock import AsyncMock

    from api.services.agent_workbench import evaluation
    from api.services.workflow import text_chat_runner

    suite = {
        "name": "greeting",
        "cases": [
            {
                "name": "greeting",
                "user_messages": ["hello"],
                "assertions": [{"kind": "contains", "value": "hello"}],
            }
        ],
    }

    async def get(_org, _sid, table, _id):
        return (
            {"snapshot": {"workflow_json": {"nodes": []}}, "fingerprint": "candidate"}
            if table == "candidates"
            else {"payload": suite, "fingerprint": "suite"}
        )

    monkeypatch.setattr(evaluation.db_client, "wb_get", get)
    for method, value in [
        ("wb_context", {"session": {"workflow_id": 1}}),
        ("wb_evaluation", {"id": "e"}),
        ("wb_create_run", 1),
        ("wb_emit", None),
        ("wb_heartbeat", None),
        ("wb_evaluation_result", None),
        ("update_workflow_run", None),
    ]:
        monkeypatch.setattr(evaluation.db_client, method, AsyncMock(return_value=value))
    responses = [
        SimpleNamespace(
            checkpoint={},
            gathered_context={},
            is_completed=False,
            assistant_text="hello",
            events=[],
            usage={},
        ),
        SimpleNamespace(
            checkpoint={},
            gathered_context={},
            is_completed=False,
            assistant_text=None,
            events=[],
            usage={},
        ),
    ]
    monkeypatch.setattr(
        text_chat_runner,
        "execute_text_chat_pending_turn",
        AsyncMock(side_effect=responses),
    )
    result = await evaluation.evaluate(
        1,
        "s",
        {"id": "job", "lease_token": "lease"},
        "c",
        "suite",
        SimpleNamespace(llm=SimpleNamespace(model="test")),
    )
    assert result["summary"]["error"] == 1 and result["summary"]["passed"] == 0


def test_runtime_failure_keeps_provider_cause_and_redacts_credentials():
    from types import SimpleNamespace

    policy = EvaluationPolicy(
        SimulationCase(
            name="case",
            user_messages=["hello"],
            assertions=[{"kind": "contains", "value": "hello"}],
        ),
        SimpleNamespace(llm=SimpleNamespace(api_key="raw-provider-key")),
    )
    policy.record_runtime_error(
        SimpleNamespace(
            error="Error code: 400 - Arrearage: account has insufficient credits; "
            "supplied raw-provider-key; Authorization: Bearer private-token; api_key='secret-key'"
        )
    )
    assert len(policy.errors) == 1
    assert "400" in policy.errors[0] and "Arrearage" in policy.errors[0]
    assert "insufficient credits" in policy.errors[0]
    assert all(
        secret not in policy.errors[0]
        for secret in ["raw-provider-key", "private-token", "secret-key"]
    )
    policy.record_runtime_error(SimpleNamespace(error="failure " * 2000))
    assert len(policy.errors[-1]) < 1050
