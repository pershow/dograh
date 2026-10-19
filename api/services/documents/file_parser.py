"""Standalone parser subprocess. No credentials, HTTP requests or macro execution."""

import base64
import csv
import io
import json
import sys
import zipfile
from pathlib import Path

MAX_TEXT = 200_000


def parse_file(path):
    suffix = path.suffix.lower()
    parts = []
    if suffix in {".docx", ".xlsx"}:
        with zipfile.ZipFile(path) as archive:
            entries = archive.infolist()
            if (
                len(entries) > 5000
                or sum(e.file_size for e in entries) > 80 * 1024 * 1024
            ):
                raise ValueError("Office document exceeds expanded-size limit")
            if any(
                "vbaproject" in e.filename.lower() or e.flag_bits & 1 for e in entries
            ):
                raise ValueError("Macros and encrypted Office files are unsupported")
    if suffix in {".txt", ".md", ".csv", ".json"}:
        text = path.read_text(encoding="utf-8-sig")
        if "\0" in text:
            raise ValueError("Binary data is not a text document")
        if suffix == ".json":
            text = json.dumps(json.loads(text), ensure_ascii=False, indent=2)
        elif suffix == ".csv":
            rows = csv.reader(io.StringIO(text))
            rows = [row for _, row in zip(range(5001), rows)]
            if len(rows) > 5000:
                raise ValueError("CSV exceeds 5000-row limit")
            text = "\n".join(" | ".join(row) for row in rows)
        parts = [("document", text)]
    elif suffix == ".pdf":
        from pypdf import PdfReader

        reader = PdfReader(path)
        if reader.is_encrypted or len(reader.pages) > 200:
            raise ValueError("PDF must be unencrypted and at most 200 pages")
        parts = [
            (f"page {i + 1}", page.extract_text() or "")
            for i, page in enumerate(reader.pages)
        ]
        if not any(text.strip() for _, text in parts):
            raise ValueError(
                "PDF has no extractable text; upload a text export or page images"
            )
    elif suffix == ".docx":
        from docx import Document

        doc = Document(path)
        parts = [("paragraphs", "\n".join(p.text for p in doc.paragraphs))]
        parts += [
            (
                f"table {i + 1}",
                "\n".join(" | ".join(c.text for c in r.cells) for r in table.rows),
            )
            for i, table in enumerate(doc.tables)
        ]
    elif suffix == ".xlsx":
        from openpyxl import load_workbook

        book = load_workbook(path, read_only=True, data_only=True, keep_links=False)
        try:
            count = 0
            for sheet in book:
                lines = []
                for row in sheet.iter_rows(values_only=True):
                    count += len(row)
                    if count > 50000:
                        raise ValueError("Spreadsheet exceeds 50000-cell limit")
                    lines.append(" | ".join("" if v is None else str(v) for v in row))
                parts.append((sheet.title, "\n".join(lines)))
        finally:
            book.close()
    elif suffix in {".png", ".jpg", ".jpeg", ".webp"}:
        from PIL import Image

        Image.MAX_IMAGE_PIXELS = 20_000_000
        with Image.open(path) as image:
            width, height = image.size
            if width * height > 20_000_000 or image.format not in {
                "PNG",
                "JPEG",
                "WEBP",
            }:
                raise ValueError("Unsupported or oversized image")
            image.verify()
        with Image.open(path) as source:
            image = source.convert("RGB")
            image.thumbnail((768, 768))
            output = io.BytesIO()
            image.save(output, "JPEG", quality=70)
        return {
            "kind": "image",
            "width": width,
            "height": height,
            "text": "",
            "segments": [],
            "parser": "pillow",
            "model_image": "data:image/jpeg;base64,"
            + base64.b64encode(output.getvalue()).decode(),
        }
    else:
        raise ValueError("Unsupported file format")
    segments, remaining = [], MAX_TEXT
    for location, text in parts:
        if remaining <= 0:
            break
        segments.append({"location": location, "text": text[:remaining]})
        remaining -= len(text)
    return {
        "kind": "text",
        "text": "\n\n".join(s["text"] for s in segments),
        "segments": segments,
        "truncated": sum(len(t) for _, t in parts) > MAX_TEXT,
        "parser": "local-v1",
    }


def main():
    import resource

    resource.setrlimit(resource.RLIMIT_CPU, (20, 20))
    if sys.platform == "linux":
        resource.setrlimit(resource.RLIMIT_AS, (768 * 1024 * 1024, 768 * 1024 * 1024))
    try:
        result = parse_file(Path(sys.argv[1]))
    except (ValueError, UnicodeError, zipfile.BadZipFile) as error:
        result = {"error": str(error)}
    print(json.dumps(result, ensure_ascii=False))


if __name__ == "__main__":
    main()
